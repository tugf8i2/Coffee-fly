import unittest
from io import BytesIO
from openpyxl import load_workbook
from app.services.exportacion_reportes import crear_excel, crear_pdf


class ReportExportTests(unittest.TestCase):
    periodo = {'desde': '2026-09-01', 'hasta': '2026-09-16'}

    def test_excel_preserves_numbers_and_literal_names(self):
        content = crear_excel(self.periodo, [('Café por caficultor', ['Caficultor', 'Kilogramos'], [['=Nombre', 1234.5]])])
        sheet = load_workbook(BytesIO(content)).active
        self.assertEqual(sheet['A5'].value, '=Nombre')
        self.assertEqual(sheet['A5'].data_type, 's')
        self.assertEqual(sheet['B5'].value, 1234.5)
        self.assertEqual(sheet.freeze_panes, 'A5')
        self.assertEqual(sheet.auto_filter.ref, 'A4:B5')

    def test_empty_exports_remain_readable(self):
        sections = [('Resumen diario', ['Fecha', 'Recolecciones', 'Kilogramos'], [])]
        sheet = load_workbook(BytesIO(crear_excel(self.periodo, sections))).active
        self.assertEqual(sheet['A5'].value, 'Sin datos para el período')
        self.assertTrue(crear_pdf(self.periodo, sections).startswith(b'%PDF'))

    def test_pdf_handles_long_names_special_characters_and_many_rows(self):
        rows = [[f'{i} María & José <Finca> ' + 'Apellido ' * 15, 1250.75] for i in range(120)]
        content = crear_pdf(self.periodo, [('Café por caficultor', ['Caficultor', 'Kilogramos'], rows)])
        self.assertTrue(content.startswith(b'%PDF'))
        self.assertGreater(content.count(b'/Type /Page'), 2)


if __name__ == '__main__':
    unittest.main()
