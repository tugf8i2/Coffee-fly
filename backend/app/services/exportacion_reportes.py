"""Presentación de reportes; recibe las mismas filas que la consulta JSON."""
from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle


def crear_excel(periodo, secciones, nota='Recolecciones no canceladas'):
    libro = Workbook()
    libro.remove(libro.active)
    for titulo, columnas, filas in secciones:
        hoja = libro.create_sheet(titulo[:31])
        hoja.append(['Coffee Fly · Reporte operativo'])
        hoja.append([f"Período: {periodo['desde']} a {periodo['hasta']}"])
        hoja.append([nota])
        hoja.append(columnas)
        for fila in filas:
            hoja.append(fila)
            for cell in hoja[hoja.max_row]:
                # Los nombres se conservan como texto, incluso si comienzan con '='.
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        if not filas:
            hoja.append(['Sin datos para el período'])
        for row in hoja:
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if cell.row > 4 and isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0.00' if columnas[cell.column - 1] == 'Kilogramos' else '#,##0'
        for cell in hoja[4]:
            cell.fill = PatternFill('solid', fgColor='123F34')
            cell.font = Font(color='FFFFFF', bold=True)
        hoja['A1'].font = Font(size=16, bold=True, color='123F34')
        hoja.column_dimensions['A'].width = 52
        for index in range(2, len(columnas) + 1):
            hoja.column_dimensions[hoja.cell(4, index).column_letter].width = 20
        hoja.freeze_panes = 'A5'
        hoja.auto_filter.ref = f'A4:{hoja.cell(max(4, hoja.max_row), len(columnas)).coordinate}'
        hoja.print_title_rows = '1:4'
        hoja.sheet_properties.pageSetUpPr.fitToPage = True
        hoja.page_setup.orientation = 'landscape'
        hoja.page_setup.paperSize = hoja.PAPERSIZE_A4
        hoja.page_setup.fitToWidth = 1
        hoja.page_setup.fitToHeight = 0
    output = BytesIO()
    libro.save(output)
    return output.getvalue()


def crear_pdf(periodo, secciones, nota='Recolecciones no canceladas · Peso en kilogramos'):
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=44,
                            title='Coffee Fly - Reporte operativo', author='Coffee Fly')
    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle('ReportCell', fontName='Helvetica', fontSize=9, leading=13)
    header_style = ParagraphStyle('ReportHeader', parent=cell_style, textColor=colors.white, fontName='Helvetica-Bold')
    story = [Paragraph('Coffee Fly · Reporte operativo', styles['Title']),
             Paragraph(f"Período: {periodo['desde']} a {periodo['hasta']}", styles['Normal']),
             Paragraph(escape(nota), styles['Normal']), Spacer(1, 16)]
    for titulo, columnas, filas in secciones:
        story.append(Paragraph(escape(titulo), styles['Heading2']))
        rows = [[Paragraph(escape(col), header_style) for col in columnas]]
        for fila in filas:
            rows.append([Paragraph(escape(f'{value:,.2f}' if columnas[i] == 'Kilogramos' else str(value)), cell_style) for i, value in enumerate(fila)])
        if not filas:
            rows.append([Paragraph('Sin datos para el período', cell_style)] + [''] * (len(columnas) - 1))
        widths = [doc.width * .65, doc.width * .35] if len(columnas) == 2 else [doc.width * .5, doc.width * .22, doc.width * .28]
        table = Table(rows, colWidths=widths, repeatRows=1, hAlign='LEFT')
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#123F34')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F0F5EF')]),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('LEFTPADDING', (0, 0), (-1, -1), 9), ('RIGHTPADDING', (0, 0), (-1, -1), 9),
            ('TOPPADDING', (0, 0), (-1, -1), 7), ('BOTTOMPADDING', (0, 0), (-1, -1), 7),
        ]))
        story.extend([table, Spacer(1, 14)])
    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont('Helvetica', 8)
        canvas.setFillColor(colors.HexColor('#526451'))
        canvas.drawString(40, 24, 'Coffee Fly · Reporte de operaciones')
        canvas.drawRightString(letter[0] - 40, 24, f'Página {document.page}')
        canvas.restoreState()
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()
